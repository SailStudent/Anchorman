from pptx.util import Inches

from static.emojis._EMOJI_REFERENCES import _EMOJI_REFERENCES

def add_emoji_to_table(emoji, slide, table, row_idx, col_idx):
    if emoji in _EMOJI_REFERENCES:
        img_path = _EMOJI_REFERENCES[emoji]
        cell = table.cell(row_idx, col_idx)
        row_heights = [r.height for r in table.rows]
        col_widths = [c.width for c in table.columns]
        
        cell_width = col_widths[col_idx]
        cell_height = row_heights[row_idx]

        img = slide.shapes.add_picture(img_path, 0, 0, width=Inches(0.4), height=Inches(0.4))

        table_left = table._graphic_frame.left
        table_top = table._graphic_frame.top
        
        cell_left = table_left + sum(col_widths[:col_idx])
        cell_top = table_top + sum(row_heights[:row_idx])

        img.left = int(cell_left + (cell_width - img.width) / 2)
        img.top = int(cell_top + (cell_height - img.height) / 2)
        
        cell.text = ""